import openpyxl as op

class Student:
    def __init__(self, name, USN, CGPA, branch, choices):
        self.name = name
        self.USN = USN
        self.CGPA = CGPA
        self.branch = branch
        self.choices=choices
        self.alloc = ""

    def __str__(self):
        return(f"Name : {self.name} , USN : {self.USN} , CGPA : {self.CGPA} , branch : {self.branch} , choices : {self.choices} , allocated elective : {self.alloc}")

    def tuplestr(self):
        return(f"{self.name},{self.USN},{self.CGPA},{self.branch},{self.alloc},{(',').join(self.choices)}")     

    def tuplestr2(self):
        return(f"{self.name},{self.USN},{self.CGPA},{self.branch}")


class Elective:
    def __init__(self,courseid,ename,branch,MaxCap=3):
        self.Name = ename
        self.courseID = courseid
        self.Branch = branch
        self.MaxCap = MaxCap
        self.NoStuds = 0
        self.MinCgpa = 0
        self.lastS = 0

    def __str__(self):
        return(f"Name : {self.Name} , ID : {self.courseID} , branch : {self.Branch}  , No. of students: {self.NoStuds}, MaxCap : {self.MaxCap}, Min CGPA : {self.MinCgpa}")

def createElectives():
    courseCode= ""
    for i in range(2,6):
        for j in range (1,4):
            val = courses.cell(i,j).value
            if j ==1 :
                courseCode = val
            elif j ==2 :
                ename = val
            elif j ==3 :
                branch = val    
        e = Elective(courseCode,ename,branch)
        Course[courseCode]=e
        
def createStud():
    for i in range(2,22):
        choice=[]
        for j in range(1,8):
            val=studs.cell(i,j).value
            if j ==1 :
                sname = val
            elif j ==2 :
                usn = val
            elif j ==3 :
                cgpa = val 
            elif j ==4 :
                sbranch = val
            else:
                choice.append(val)
        s=Student(sname,usn,cgpa,sbranch,choice)
        Stud.append(s)        

def RawResult(Studs):
    newWorkbook = op.load_workbook("Allotment.xlsx")
    sheet=newWorkbook.active
    l=[]
    t1=("Student Name","USN","CGPA","Branch","Allotted Course","Choice 1","Choice 2","Choice 3")
    l.append(t1)
    for i in Studs:
        v=i.tuplestr().split(',')
        t=tuple(v)
        l.append(t)
    for j in l:
        sheet.append(j)
        
    newWorkbook.save(filepath)   

def writeCourses(sheet_names,Stud):

    newWorkbook = op.load_workbook("Allotment.xlsx")
    for i in sheet_names:
        newWorkbook.active = newWorkbook[i]
        print(newWorkbook.active.title)
        t1=("Student Name","USN","CGPA","Branch")
        l=[]
        l.append(t1)
        for j in Stud:
            if(j.alloc==i):
                u=j.tuplestr2().split(',')
                t1=tuple(u)
                l.append(t1)
        for j in l:
            newWorkbook.active.append(j)
    newWorkbook.save(filepath)
                                   
def AllotCourse(Stud,i):

    for j in Stud.choices:
        if Course[j].NoStuds < Course[j].MaxCap:
            Stud.alloc=j
            Course[j].prevS=Stud
            Course[j].NoStuds+=1
            Course[j].MinCgpa=Stud.CGPA
            Course[j].lastS=i
            return -1
        
        elif Course[j].NoStuds < (Course[j].MaxCap + buffer) :
            if Stud.CGPA == Course[j].MinCgpa:
                Stud.alloc=j
                Course[j].prevS=Stud
                Course[j].NoStuds+=1
                Course[j].lastS=i
                return -1

        elif  Course[j].NoStuds == (Course[j].MaxCap + buffer) :
            if Stud.CGPA == Course[j].MinCgpa and Stud.choices.index(j)<Course[j].prevS.choices.index(j):
                Stud.alloc=j
                Course[j].prevS=Stud                    
                prev=Course[j].lastS                  
                Course[j].lastS=i
                return prev
            
    return -1         
            
filepath = r"D:/ABHIJITH IYER R/jce/notes , textbooks etc/sem 5/open elective allotment project/Allotment.xlsx"
#r"d:/CSE/5TH SEMESTER/Open Elective Allotment/Allotment.xlsx"
# r"D:/ABHIJITH IYER R/jce/notes , textbooks etc/sem 5/open elective allotment project/Allotment.xlsx"
if __name__ == "__main__":

    wb=op.load_workbook("Students.xlsx")    
    
    studs=wb['Studs']
    courses = wb['Courses_offered']
    Stud=[]
    Course = {}
    buffer=2

    createElectives()
    createStud()
                                                          
    Stud=sorted(Stud,key = lambda Stud: Stud.CGPA,reverse=True)

    for i in range(len(Stud)):
        reallot= -1
        reallot=AllotCourse(Stud[i],i)
        while reallot > 0 :
            Stud[reallot].alloc=""
            reallot=AllotCourse(Stud[reallot],reallot)
                    
    Stud=sorted(Stud,key = lambda Stud: Stud.USN)                
                  
    print("*********************************************************")
    for i in Stud:
        print(i)
        print() 
    print("*********************************************************")
    for i in Course:
        print(Course[i])
        print()
    
    res = op.Workbook()
    sheet=res.active
    sheet.title="Allotment"
    sheet_names=list(Course.keys())
    ws=[]

    for i in range(len(sheet_names)):
        wsx=res.create_sheet(index=(i+1),title=sheet_names[i])
        ws.append(wsx)

    res.save(filepath) 
    RawResult(Stud)
    writeCourses(sheet_names,Stud)
